from flask import Flask, request, jsonify, send_file
import hmac
import hashlib
import json
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill
import anthropic
import requests

app = Flask(__name__)

# Se ejecuta apenas se importa este archivo - así funciona tanto con
# "python webhook.py" (pruebas locales) como con "gunicorn webhook:app" (Render)

# Cargar tokens desde variables de entorno (después los configuraremos en Render)
VERIFY_TOKEN = os.getenv("VERIFY_TOKEN")  # Texto que TÚ inventas, se usa solo para la verificación inicial del webhook (paso GET)
APP_SECRET = os.getenv("APP_SECRET")  # Clave secreta de la app (Configuración > Básica) - se usa para validar la firma de cada mensaje real
INSTAGRAM_TOKEN = os.getenv("INSTAGRAM_TOKEN")  # Token de acceso de Instagram
FACEBOOK_TOKEN = os.getenv("FACEBOOK_TOKEN")  # Token de acceso de Facebook
CLAUDE_API_KEY = os.getenv("CLAUDE_API_KEY")  # API Key de Anthropic
DOWNLOAD_SECRET = os.getenv("DOWNLOAD_SECRET", "clisa-descarga-2026")  # texto que TÚ inventas, para poder descargar el Excel de forma segura
EXCEL_FILE = "mensajes_clisa.xlsx"

# Palabras clave para pre-filtrado
KEYWORDS_LENTES = ["lentes", "armazón", "marco", "graduación", "receta", "óptico", "oftalmólogo", "vista", "gafas", "anteojos"]
KEYWORDS_CONSULTA = ["cita", "consulta", "horario", "abierto", "teléfono", "dirección", "agendar", "reserva", "disponibilidad"]

# Set en memoria para evitar procesar el mismo mensaje dos veces si Meta lo reintenta
# (se reinicia si el servidor se reinicia - es una protección básica, no perfecta)
_processed_message_ids = set()

def is_duplicate(message_id):
    if message_id in _processed_message_ids:
        return True
    _processed_message_ids.add(message_id)
    return False

def init_excel():
    """Crea el archivo Excel si no existe"""
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mensajes"
        
        # Headers
        headers = ["Fecha", "Plataforma", "Remitente", "ID Remitente", "Mensaje", "Categoría", "Link a Conversación"]
        ws.append(headers)
        
        # Estilos para headers
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        
        # Ancho de columnas
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 40
        ws.column_dimensions['F'].width = 15
        ws.column_dimensions['G'].width = 50
        
        wb.save(EXCEL_FILE)

def verify_webhook_signature(body, signature):
    """Verifica que el mensaje venga realmente de Meta (usa la Clave secreta de la app, NO el verify token)"""
    expected_signature = hmac.new(
        APP_SECRET.encode(),
        body,
        hashlib.sha256
    ).hexdigest()
    return hmac.compare_digest(signature, expected_signature)

def classify_message(message_text):
    """Clasifica el mensaje usando Claude API"""
    try:
        client = anthropic.Anthropic(api_key=CLAUDE_API_KEY)
        
        prompt = f"""Clasifica el siguiente mensaje en UNA de estas categorías:
        
1. "solicitud_lentes": Si la persona pregunta o pide sobre lentes, armazones, graduación, recetas ópticas, etc.
2. "solicitud_consulta": Si la persona solicita cita, consulta, información de horarios, direcciones, etc.
3. "otro": Si no encaja en las categorías anteriores.

Responde SOLO con la categoría, sin explicaciones. Ejemplo: "solicitud_lentes"

Mensaje: "{message_text}"
"""
        
        message = client.messages.create(
            model="claude-haiku-4-5-20251001",
            max_tokens=50,
            messages=[
                {"role": "user", "content": prompt}
            ]
        )
        
        category = message.content[0].text.strip().lower()
        
        # Validar que sea una categoría válida
        if "solicitud_lentes" in category:
            return "Solicitud de Lentes"
        elif "solicitud_consulta" in category:
            return "Solicitud de Consulta"
        else:
            return "Otro"
    
    except Exception as e:
        print(f"Error al clasificar con Claude: {e}")
        # Fallback: pre-filtrado simple
        message_lower = message_text.lower()
        if any(kw in message_lower for kw in KEYWORDS_LENTES):
            return "Solicitud de Lentes"
        elif any(kw in message_lower for kw in KEYWORDS_CONSULTA):
            return "Solicitud de Consulta"
        else:
            return "Otro"

def add_to_excel(fecha, plataforma, remitente, id_remitente, mensaje, categoria, link):
    """Agrega un mensaje al Excel"""
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        
        ws.append([fecha, plataforma, remitente, id_remitente, mensaje, categoria, link])
        wb.save(EXCEL_FILE)
    except Exception as e:
        print(f"Error al escribir en Excel: {e}")

def get_sender_name_facebook(sender_id):
    """Obtiene el nombre del remitente desde Facebook"""
    try:
        url = f"https://graph.facebook.com/v18.0/{sender_id}"
        params = {"fields": "name", "access_token": FACEBOOK_TOKEN}
        response = requests.get(url, params=params)
        data = response.json()
        return data.get("name", f"Usuario {sender_id}")
    except:
        return f"Usuario {sender_id}"

def get_sender_name_instagram(sender_id):
    """Obtiene el nombre del remitente desde Instagram"""
    try:
        url = f"https://graph.instagram.com/v18.0/{sender_id}"
        params = {"fields": "username", "access_token": INSTAGRAM_TOKEN}
        response = requests.get(url, params=params)
        data = response.json()
        return data.get("username", f"Usuario {sender_id}")
    except:
        return f"Usuario {sender_id}"

init_excel()  # crea el archivo Excel si no existe, corre siempre al cargar el módulo

@app.route("/webhook", methods=["GET"])
def webhook_get():
    """Verifica el webhook con Meta"""
    mode = request.args.get("hub.mode")
    token = request.args.get("hub.verify_token")
    challenge = request.args.get("hub.challenge")
    
    if mode == "subscribe" and token == VERIFY_TOKEN:
        return challenge, 200
    else:
        return "Forbidden", 403

@app.route("/webhook", methods=["POST"])
def webhook_post():
    """Recibe y procesa los mensajes de Meta"""
    
    # Verificar firma
    signature = request.headers.get("X-Hub-Signature-256")
    if not signature or not verify_webhook_signature(request.data, signature.replace("sha256=", "")):
        return "Unauthorized", 401
    
    data = request.get_json()
    plataforma_obj = data.get("object")  # "page" = Facebook, "instagram" = Instagram

    if plataforma_obj not in ("page", "instagram"):
        # Evento que no nos interesa (Meta manda otros tipos de object)
        return jsonify({"status": "ignored"}), 200

    plataforma = "Facebook" if plataforma_obj == "page" else "Instagram"

    for entry in data.get("entry", []):
        for messaging_event in entry.get("messaging", []):
            message_obj = messaging_event.get("message")

            # Ignorar eco de mensajes que la propia página/cuenta envió (respuestas del equipo)
            if not message_obj or message_obj.get("is_echo"):
                continue

            message_id = message_obj.get("mid")
            message_text = message_obj.get("text", "")
            if not message_text:
                continue  # imagen/sticker/audio sin texto - se ignora por ahora

            sender_id = messaging_event["sender"]["id"]
            timestamp = messaging_event.get("timestamp", int(datetime.now().timestamp() * 1000))

            # Evitar procesar el mismo mensaje dos veces si Meta lo reenvía
            if message_id and is_duplicate(message_id):
                continue

            if plataforma == "Facebook":
                sender_name = get_sender_name_facebook(sender_id)
                link = f"https://www.facebook.com/messages/t/{sender_id}"
            else:
                sender_name = get_sender_name_instagram(sender_id)
                link = f"https://www.instagram.com/direct/t/{sender_id}"

            categoria = classify_message(message_text)
            fecha = datetime.fromtimestamp(timestamp / 1000).strftime("%Y-%m-%d %H:%M")

            add_to_excel(fecha, plataforma, sender_name, sender_id, message_text, categoria, link)
            print(f"[{plataforma}] {sender_name} ({sender_id}): {message_text} → {categoria}")

    return jsonify({"status": "ok"}), 200

@app.route("/health", methods=["GET"])
def health():
    """Health check para Render"""
    return jsonify({"status": "healthy"}), 200

@app.route("/descargar-excel", methods=["GET"])
def descargar_excel():
    """
    Descarga el archivo Excel actual del servidor.
    Solo funciona si se manda el parámetro correcto ?clave=TU_DOWNLOAD_SECRET
    para que no cualquiera pueda descargarlo con solo saber la URL.
    Ejemplo: https://clisa-webhook.onrender.com/descargar-excel?clave=clisa-descarga-2026
    """
    clave = request.args.get("clave")
    if clave != DOWNLOAD_SECRET:
        return "No autorizado", 403

    if not os.path.exists(EXCEL_FILE):
        return "El archivo Excel todavía no existe en el servidor", 404

    return send_file(EXCEL_FILE, as_attachment=True, download_name="mensajes_clisa.xlsx")

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", 5000)), debug=False)
